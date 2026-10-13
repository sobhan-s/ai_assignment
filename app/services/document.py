"""
Document Service
Handles parsing raw files and chunking them.
"""

from __future__ import annotations

import io
from pathlib import Path

import docx
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from pypdf import PdfReader

from app.constants import SUPPORTED_EXTENSIONS
from app.exceptions import EmptyDocumentError, UnsupportedFileTypeError


class DocumentService:
    def __init__(
        self,
        *,
        chunk_size: int = 1_000,
        chunk_overlap: int = 200,
    ) -> None:
        self._splitter = RecursiveCharacterTextSplitter(
            chunk_size=chunk_size,
            chunk_overlap=chunk_overlap,
            separators=["\n\n", "\n", " ", ""],
        )

    def parse(self, filename: str, content: bytes) -> str:
        ext = Path(filename).suffix.lower()

        if ext not in SUPPORTED_EXTENSIONS:
            raise UnsupportedFileTypeError(
                f"Unsupported file type '{ext}'. Allowed: {sorted(SUPPORTED_EXTENSIONS)}"
            )

        if ext == ".pdf":
            reader = PdfReader(io.BytesIO(content))
            text = "\n".join(
                page.extract_text() for page in reader.pages if page.extract_text()
            )
        elif ext == ".docx":
            doc = docx.Document(io.BytesIO(content))
            text = "\n".join(p.text for p in doc.paragraphs if p.text)
        else:
            text = content.decode("utf-8", errors="replace")

        if not text.strip():
            raise EmptyDocumentError(f"File '{filename}' contains no readable text.")

        return text

    def chunk(self, text: str, *, source: str) -> list[Document]:
        return [
            Document(page_content=chunk, metadata={"source": source})
            for chunk in self._splitter.split_text(text)
        ]


default_document_service = DocumentService()
